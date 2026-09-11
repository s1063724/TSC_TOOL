#!/usr/bin/env python3
"""Generate MCS dispatch CSV from a Workstation xlsx file.

Pair rule: within each pool (defined by Zone), enumerate all ordered
(Source, Dest) permutations where Source != Dest. Enabled=False rows and
devices in the exclude list are skipped.

Field defaults follow the on-site dispatch format:
  Group,Trigger Time,CarrierID,Lot ID,Lot Num,Carrier Type,
  Source,Dest,Priority,Replace,Back,BackCarrierID,Back Carrier Type,Execute Time

Trigger Time increments every N rows by M minutes (default 10/5), wrapping
at 24 hours. Carrier Type per pool cycles through a list.
"""
import argparse
import csv
import itertools
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.stderr.write("openpyxl not installed. Run: pip install --user openpyxl\n")
    sys.exit(1)

DEFAULT_POOLS = {
    "MGZ_POOL": {
        "zones": ["zone_5F_AMR_MGZ", "zone_5F_AMR_MGZ_OVEN", "zone_5F_AMR_MGZ_WB"],
        "carrier_types": ["MAG1"],
    },
    "CST_POOL": {
        "zones": ["zone_5F_AMR_CST"],
        "carrier_types": ["F08", "F12", "FOUP08", "FOUP12"],
    },
}

DEFAULT_EXCLUDE = [
    "MGZ_NG_PORT_02", "MGZ_NG_PORT_03", "MGZ_NG_PORT_04", "CST_NG_PORT_02",
]

HEADER = [
    "Group", "Trigger Time", "CarrierID", "Lot ID", "Lot Num",
    "Carrier Type", "Source", "Dest", "Priority", "Replace",
    "Back", "BackCarrierID", "Back Carrier Type", "Execute Time",
]

FIELD_DEFAULTS = {
    "Group": "*", "CarrierID": "GY001", "Lot ID": " ", "Lot Num": "0",
    "Priority": "0", "Replace": "0", "Back": "*",
    "BackCarrierID": " ", "Back Carrier Type": " ", "Execute Time": "0",
}


def fmt_time(idx, rows_per_step, min_step, wrap):
    total_min = (idx // rows_per_step) * min_step
    h, m = divmod(total_min, 60)
    if wrap:
        h %= 24
    return "{:02d}:{:02d}:00".format(h, m)


def load_devices(xlsx_path, sheet_name, pools, exclude):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]

    zone_to_pool = {}
    for pool_name, info in pools.items():
        for z in info["zones"]:
            zone_to_pool[z] = pool_name

    devices_by_pool = {name: [] for name in pools}
    exclude_set = set(exclude)
    for row in ws.iter_rows(min_row=2, values_only=True):
        device_id, zone_id, enabled = row[3], row[6], row[20]
        if not enabled or not device_id:
            continue
        if device_id in exclude_set:
            continue
        pool_name = zone_to_pool.get(zone_id)
        if pool_name:
            devices_by_pool[pool_name].append(device_id)
    return devices_by_pool


def build_rows(devices_by_pool, pools, opts):
    rows = []
    idx = 0
    for pool_name, devs in devices_by_pool.items():
        if len(devs) < 2:
            continue
        ctypes = pools[pool_name]["carrier_types"] or ["C12"]
        for src, dst in itertools.permutations(devs, 2):
            ct = ctypes[idx % len(ctypes)]
            cid = opts["carrier_id"]
            if opts["carrier_id_incr"]:
                cid = _increment_cid(opts["carrier_id"], idx)
            rows.append([
                FIELD_DEFAULTS["Group"],
                fmt_time(idx, opts["rows_per_step"], opts["min_step"], opts["wrap"]),
                cid,
                FIELD_DEFAULTS["Lot ID"],
                FIELD_DEFAULTS["Lot Num"],
                ct, src, dst,
                FIELD_DEFAULTS["Priority"],
                FIELD_DEFAULTS["Replace"],
                FIELD_DEFAULTS["Back"],
                FIELD_DEFAULTS["BackCarrierID"],
                FIELD_DEFAULTS["Back Carrier Type"],
                FIELD_DEFAULTS["Execute Time"],
            ])
            idx += 1
            if opts["max_rows"] and idx >= opts["max_rows"]:
                return rows
    return rows


def _increment_cid(base, offset):
    import re
    m = re.match(r"^(.*?)(\d+)$", str(base))
    if not m:
        return "{}{}".format(base, offset)
    prefix, num = m.group(1), m.group(2)
    return "{}{}".format(prefix, str(int(num) + offset).zfill(len(num)))


def main():
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("input", nargs="?", default="examples/Workstation_5F.xlsx",
                   help="Workstation xlsx path (default: examples/Workstation_5F.xlsx)")
    p.add_argument("-o", "--output", default="output/dispatch.csv",
                   help="Output CSV path (default: output/dispatch.csv)")
    p.add_argument("--sheet", default="Workstation", help="Sheet name")
    p.add_argument("--pools-json", help="Path to JSON file overriding pool definitions")
    p.add_argument("--exclude", default=",".join(DEFAULT_EXCLUDE),
                   help="Comma-separated device IDs to exclude")
    p.add_argument("--rows-per-step", type=int, default=10)
    p.add_argument("--min-step", type=int, default=5)
    p.add_argument("--no-wrap", action="store_true", help="Do not wrap hours at 24h")
    p.add_argument("--carrier-id", default="GY001")
    p.add_argument("--carrier-id-incr", action="store_true",
                   help="Increment CarrierID per row")
    p.add_argument("--max-rows", type=int, default=0, help="0 = no limit")
    args = p.parse_args()

    if args.pools_json:
        with open(args.pools_json) as f:
            pools = json.load(f)
    else:
        pools = DEFAULT_POOLS

    exclude = [s.strip() for s in args.exclude.split(",") if s.strip()]

    devices_by_pool = load_devices(args.input, args.sheet, pools, exclude)
    print("Pools:")
    total = 0
    for name, devs in devices_by_pool.items():
        print("  {}: {} devices".format(name, len(devs)))
        total += len(devs)
    print("  (total {})".format(total))

    opts = {
        "rows_per_step": args.rows_per_step,
        "min_step": args.min_step,
        "wrap": not args.no_wrap,
        "carrier_id": args.carrier_id,
        "carrier_id_incr": args.carrier_id_incr,
        "max_rows": args.max_rows,
    }
    rows = build_rows(devices_by_pool, pools, opts)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or ".", exist_ok=True)
    with open(args.output, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        w.writerows(rows)

    print("Wrote {} rows to {}".format(len(rows), args.output))
    if rows:
        print("First 3:")
        for r in rows[:3]:
            print("  " + ",".join(r))


if __name__ == "__main__":
    main()
